from pathlib import Path
import re,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
s=(R/'docs/video/Script-video-DVWA.md').read_text()
# Conserver les procédures complètes, réécrire la narration collective.
verbs={'cherche':'cherchons','compare':'comparons','complète':'complétons','conserve':'conservons','couvre':'couvrons','distingue':'distinguons','dois':'devons','démontre':'démontrons','modifie':'modifions','présente':'présentons','reconnais':'reconnaissons','regarde':'regardons','teste':'testons','vais':'allons','vérifie':'vérifions','connais':'connaissons','marque':'marquons','montre':'montrons','précise':'précisons','rejoue':'rejouons','relie':'relions','suis':'suivons'}
for a,b in verbs.items():
 s=re.sub(r'\bJe '+a+r'\b','Nous '+b,s); s=re.sub(r'\bje '+a+r'\b','nous '+b,s)
for a,b in {'Je ne présente':'Nous ne présentons','je ne peux':'nous ne pouvons','Je n’ai':'Nous n’avons','je n’ai':'nous n’avons','J’ai':'Nous avons','j’ai':'nous avons','J’observe':'Nous observons','j’observe':'nous observons','J’utilise':'Nous utilisons','j’examine':'nous examinons','Mon objectif':'Notre objectif','mon travail':'notre travail','mes propres':'nos propres','mes instances':'nos instances','mon test':'notre test','ta vidéo':'notre vidéo','[ton nom]':'Groupe 12'}.items(): s=s.replace(a,b)
s=s.replace('/home/modou/dvwa-study','/home/modou/groupe12')
start=s.index('## 01.'); end=s.index('## Parcours systématique')
chapters=s[start:end]
entries=json.loads((R/'evidence/attacks/results.json').read_text())
intro='''# Script vidéo complet — Notre étude DVWA
GROUPE 12 • Version V2 • 15 septembre 2026
Du « Bonjour » à la conclusion • 19 modules et tous les sous-tests du rapport

## Mode d’emploi
Lire à voix haute uniquement les passages entre guillemets dans les rubriques « À dire ». Les rubriques « À l’écran » et « À filmer et à faire » indiquent les gestes à effectuer. Faire une pause pendant les manipulations et laisser le résultat visible quelques secondes.
Les résultats de référence viennent des essais enregistrés le 14 septembre 2026. Lors du nouvel enregistrement, annoncer ce qui se produit réellement. Une capture ancienne doit être présentée comme une archive datée. Les charges de démonstration du guide utilisent parfois des marqueurs VIDEO différents des marqueurs G12 des archives : elles doivent être vérifiées pendant la nouvelle prise.
Notre parcours principal couvre les 22 sous-tests enregistrés. Les comparaisons de niveaux et les sous-exercices supplémentaires sont décrits pour permettre un tournage complet ; ils ne sont pas tous déjà validés. Le niveau impossible désigne une protection fournie par DVWA et ne constitue pas notre correctif personnel.
Durée indicative : 70 à 100 minutes pour les 19 modules, davantage si toutes les comparaisons sont exécutées. Nous pouvons enregistrer par chapitres puis assembler les prises. Le PowerPoint de 28 diapositives sert de support ; cette vidéo n’est pas limitée à 30 pages ou minutes.

## 0. Préparer les fenêtres avant de commencer — ne pas lire
1. Ouvrir le PowerPoint Presentation-DVWA-V2.pptx et le rapport V2.
2. Ouvrir deux profils de navigateur : référence sur http://127.0.0.1:4280 et copie de travail sur http://127.0.0.1:4281. Les cookies de sécurité peuvent être partagés entre les ports dans un même profil.
3. Préparer le terminal dans /home/modou/groupe12. Démarrer les services avec docker compose up -d, puis contrôler docker compose ps. Utiliser python3 scripts/check.py pour le contrôle sans réinitialisation ; ne pas ajouter --init pour une simple reprise.
4. Préparer les comptes fictifs admin / password et gordonb / abc123. Vérifier les identifiants avant la prise.
5. Ouvrir les outils développeur, onglets Réseau, Stockage et Console. Pour les requêtes modifiées, préparer Burp Repeater avec la session de l’instance concernée ; utiliser les jetons frais quand le formulaire en demande.
6. Préparer les fichiers dans docs/video/fixtures. Les commandes Python de ce script sont à lancer depuis /home/modou/groupe12 ; les chemins fixtures/ désignent docs/video/fixtures/.
7. Pour CSRF et redirection, lancer dans un terminal séparé :
```bash
cd /home/modou/groupe12
python3 -m http.server 4290 --bind 127.0.0.1 --directory docs/video/fixtures
```
8. Noter les changements de mot de passe, de cookie ou de données, puis remettre l’état initial après chaque scène. Les démonstrations CSRF et CAPTCHA du guide utilisent Video-Lab-2026! ; les archives du rapport utilisent G12-Test-2026!.
9. Pour chaque prise, conserver l’identifiant du test, le compte, le niveau, l’entrée et l’effet. Afficher le niveau avant de commencer. Le port 4281 ne signifie pas que l’application est corrigée.
10. Ne lancer aucun script de tous les tests pendant la narration : effectuer chaque démonstration au rythme de l’explication et vérifier son résultat.

## 1. Bonjour et présentation — à dire
« Bonjour à toutes et à tous. Nous sommes le groupe 12. Dans cette vidéo, nous présentons notre projet d’étude des attaques web et de leurs contre-mesures avec DVWA, Damn Vulnerable Web Application.
Notre objectif est de comprendre comment une entrée utilisateur peut provoquer une requête SQL détournée, une exécution de script, une modification non autorisée ou une divulgation de données.
Nous allons d’abord présenter notre environnement. Ensuite, nous parcourrons les dix-neuf modules, en montrant les manipulations et les résultats. Nous terminerons par nos limites et les étapes nécessaires pour valider nos correctifs.
DVWA est une application pédagogique développée par ses auteurs. Notre travail porte sur la préparation du laboratoire, les essais, la conservation des preuves et leur analyse. »
À l’écran : diapositive 1, puis diapositive 2. Marquer une pause avant le terminal.

## 2. Présenter et démarrer notre laboratoire
À l’écran : diapositive 3, compose.yaml, puis le terminal.
À dire : « Nous utilisons cinq services : deux applications DVWA, deux bases MariaDB et un proxy Nginx. La référence est accessible sur le port 4280 ; notre copie de travail est sur le port 4281. Chaque instance possède son réseau et ses volumes. Nous conservons ainsi les données entre deux reprises. »
À exécuter :
```bash
cd /home/modou/groupe12
docker compose up -d
docker compose ps
```
À dire si les états sont corrects : « Nous voyons les cinq services démarrés et les deux bases à l’état Healthy. Nous avons démarré le laboratoire en conservant les données existantes. »
À dire en cas d’écart : « Nous rencontrons un problème de démarrage. Nous vérifions les journaux avant de poursuivre les essais. »
À l’écran : ouvrir les deux pages de connexion, saisir admin / password, montrer l’accueil, puis DVWA Security et low. Présenter Setup Check sans cliquer sur Create / Reset Database.
À dire : « Nous vérifions maintenant l’accès aux deux interfaces. Notre copie de travail reste vulnérable à ce stade ; son nom ne constitue pas une preuve de correction. »

## 3. Expliquer notre méthode et notre bilan de référence
À l’écran : diapositives 6 et 7.
À dire : « Pour chaque test, nous montrons le contexte, l’entrée soumise et le résultat. Quand cela est possible, nous comparons avec un usage normal. Nous vérifions l’état réel, car une réponse HTTP positive ne suffit pas à confirmer une attaque.
Nos journaux contiennent vingt-deux sous-tests : dix-sept effets confirmés, quatre essais à reprendre et une modification API dont la persistance n’est pas confirmée. Pendant cette vidéo, nous distinguerons les résultats déjà documentés et ceux que nous devons encore vérifier.
Nous commençons par les injections SQL. »

## Plan des démonstrations
T01 SQL ; T02 SQL aveugle ; T03 commandes ; T04 XSS réfléchie ; T05 XSS stockée ; T06 XSS DOM ; T07 CSRF ; T08 inclusion ; T09 téléversement ; T10 autorisation ; T11 BAC ; T12 mots de passe ; T13 sessions ; T14 CAPTCHA ; T15 CSP ; T16 JavaScript ; T17 redirection ; T18a XOR et T18b ECB ; T19a versions API, T19b modification API et T19c commande API.
Les compléments cryptographiques high/impossible et le parcours d’authentification API sont inclus dans les sections 18 et 19, avec leurs limites.

'''
special={
'T06':'« Dans nos archives, ce test est à reprendre : la vérification de l’exécution JavaScript a échoué. Nous allons observer cette nouvelle tentative sans annoncer de succès à l’avance. »',
'T13':'« Notre précédente collecte a été interrompue par une navigation. Nous devons relever plusieurs valeurs du cookie avant de conclure sur leur prédictibilité. »',
'T18a':'« Notre capture XOR montre un message décodé, mais la vérification complète du script a échoué. Nous conservons ce sous-test à reprendre. »',
'T19b':'« La réponse de modification indique level égal à zéro, mais la relecture retourne un. Nous ne confirmons donc pas une élévation de privilèges persistante. »',
'T19c':'« Notre essai d’injection de commandes API reste à reprendre. Nous ne disposons pas de capture confirmant cet effet. »'}
parts=re.split(r'(?=^## \d{2}\.)',chapters,flags=re.M)
assembled=[]
for part in parts:
 if not part.strip(): continue
 n=int(re.search(r'^## (\d{2})\.',part).group(1)); ids=[f'T{n:02d}'] if n<18 else (['T18a','T18b'] if n==18 else ['T19a','T19b','T19c'])
 rows=[e for e in entries if e['id'] in ids]
 pos=part.index('\n### ')
 recap='\n### Notre point de départ — à dire avant de manipuler\n'
 for e in rows:
  if e['id'] in special: recap+=special[e['id']]+'\n'
  else:
   observed=e.get('observed','')
   if e['id']=='T02': observed='la condition vraie donne un utilisateur présent et HTTP 200 ; la condition fausse donne un utilisateur absent et HTTP 404.'
   recap+='« Pour '+e['id']+', nos essais enregistrés indiquent : '+observed+' Nous allons vérifier le résultat de cette prise. »\n'
  recap+='Archives disponibles : '+(', '.join(c['file'] for c in e['captures']) or 'aucune capture')+'.\n'
 recap+='Support : diapositive '+str(n+7)+' du PowerPoint. Les archives se trouvent dans docs/captures/attacks/.\n'
 part=part[:pos]+recap+part[pos:]
 part=part.replace('### Comparaison et contre-mesures — à dire','### Comparaison et contre-mesures — expliquer puis vérifier à l’écran')
 part+='\n### Si le résultat ne correspond pas — à dire\n« Nous n’obtenons pas l’effet attendu dans cette prise. Nous vérifions le niveau, la session, l’entrée et la réponse. Pour le moment, nous conservons ce résultat comme non confirmé. »\n'
 nxt={3:'les attaques XSS',6:'le changement de mot de passe par CSRF',9:'les contrôles d’autorisation',12:'les identifiants de session',14:'la politique CSP',17:'les sous-exercices de cryptographie',18:'la sécurité des API',19:'notre bilan final'}.get(n,'le test suivant')
 part+='\n### Transition — à dire\n« Nous conservons cette observation et nous remettons les données modifiées dans leur état initial. Nous passons maintenant à '+nxt+'. »\n\n'
 assembled.append(part)
ending='''## 20. Comparaisons complémentaires et correctifs — à dire
« Nous avons présenté les scénarios de notre étude. Les niveaux medium, high et impossible fournissent des éléments de comparaison propres à DVWA. Nous devons vérifier leur code et leur comportement pour chaque scénario.
Nos correctifs personnels devront être identifiés dans Git. Pour chacun, nous présenterons le problème initial, la modification, le rejeu de l’attaque et un usage légitime qui fonctionne encore. Nous ne présentons pas cette phase comme achevée sans ses preuves. »
À l’écran : diapositive 27. Si un correctif a réellement été développé depuis ce script, montrer son diff, son identifiant de commit, les requêtes avant/après et l’état final. Sinon, conserver cette séquence comme présentation du travail à venir.

## 21. Bilan final — à dire
« Nous arrivons au bilan de nos tests. Notre laboratoire permet d’étudier les dix-neuf modules dans deux instances séparées.
Dans les preuves enregistrées pour le rapport, nous avons vingt-deux sous-tests. Dix-sept présentent un effet confirmé. Quatre restent à reprendre : la XSS DOM, les identifiants de session, le sous-test XOR et l’injection de commandes API. Enfin, le test de modification API ne confirme pas une élévation persistante, car la valeur relue reste identique à sa valeur initiale.
Ces résultats nous apprennent à vérifier les données et les effets réels. Une capture d’une page, un code HTTP 200 ou une réponse de modification ne suffisent pas toujours.
Nous devons maintenant compléter les essais incomplets et valider les correctifs avec des tests comparables. »
À l’écran : diapositive 7 pour les nombres, puis 27. Si les essais de la nouvelle vidéo ont changé les résultats, annoncer les nouveaux résultats avec leurs preuves et préciser la différence avec les archives ; ne pas lire des nombres devenus inexacts.

## 22. Conclusion et remerciements — à dire
« Pour conclure, ce projet nous a permis de relier des entrées utilisateur à leurs conséquences sur une application web. Nous avons étudié les injections, les scripts exécutés dans le navigateur, les contrôles d’accès, les fichiers, les sessions, la cryptographie et les API.
Notre démarche repose sur trois points : observer, conserver les preuves et vérifier les limites de nos conclusions. La suite du travail consiste à transformer cette analyse en correctifs documentés, puis à vérifier leur efficacité sans casser les usages légitimes.
Merci d’avoir suivi cette présentation du groupe 12. Nous vous remercions pour votre attention et nous sommes disponibles pour répondre à vos questions. »
À l’écran : diapositive 28. Laisser la conclusion affichée quelques secondes, puis arrêter l’enregistrement.

## Annexe A. Phrases utiles pendant une difficulté
Session expirée : « Nous sommes redirigés vers la connexion. Nous rétablissons la session avant de reprendre ce test. »
Page lente : « Nous attendons la réponse réelle ; nous ne concluons pas à une attaque temporelle à partir de cette attente seule. »
CAPTCHA indisponible : « Les clés ou la connectivité nécessaires au parcours externe ne sont pas disponibles. Nous présentons la limite et l’analyse du code. »
Résultat contradictoire : « La réponse et l’état relu ne concordent pas. Nous retenons l’état vérifié et poursuivons le diagnostic. »
Capture d’archive : « Cette capture provient de notre essai enregistré le 14 septembre 2026. Nous la montrons comme preuve archivée, et non comme le résultat de la manipulation en direct. »
Protection observée : « Cette charge ne produit pas l’effet recherché. Nous vérifions aussi un usage normal et le code avant de conclure sur la protection. »

## Annexe B. Vérification avant de terminer la vidéo
- Les 19 modules ont un chapitre ; T18a/T18b et T19a/T19b/T19c sont distingués.
- Chaque succès annoncé possède un effet visible ou une preuve enregistrée clairement identifiée.
- Les quatre essais à reprendre et le résultat API non persistant sont signalés.
- Les comptes, niveaux et instances sont identifiables.
- Les mots de passe et états modifiés sont restaurés ; les captures sont sauvegardées.
- Les comparaisons supplémentaires sont marquées exécutées, non exécutées ou bloquées selon la prise.
- La conclusion et les nombres correspondent aux preuves présentées.

## Annexe C. Sources du script
GROUPE-12/Rapport-DVWA-V2.docx ; GROUPE-12/Presentation-DVWA-V2.pptx ; docs/video/Script-video-DVWA.md ; scripts/test_attacks.py ; evidence/attacks/results.json ; docs/captures/attacks/ ; sources locales upstream/vulnerabilities/ et fixtures du projet.
DVWA : https://github.com/digininja/DVWA. Sources locales identifiées par le commit b496a5d3de6b967410155e1b7d3e51e9d035eb22.
'''
content=(intro+''.join(assembled)+ending).replace('à les ', 'aux ').replace('à le ', 'au ')
content=content.replace('fixtures/js-tokens.py','docs/video/fixtures/js-tokens.py').replace('fixtures/crypto-xor.py','docs/video/fixtures/crypto-xor.py').replace('fixtures/dom-xss-url.txt','docs/video/fixtures/dom-xss-url.txt')
out=R/'GROUPE-12/Script-video-complet-DVWA-V2'
out.with_suffix('.md').write_text(content)
d=Document(); sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
for n in ['Title','Heading 1','Heading 2']:
 d.styles[n].font.color.rgb=RGBColor.from_string('14283F')
sec.header.paragraphs[0].text='GROUPE 12 • Script vidéo complet • DVWA V2'
f=sec.footer.paragraphs[0]; f.text='Script de tournage • Page '; fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
in_code=False
for line in content.splitlines():
 if line.startswith('```'): in_code=not in_code;continue
 if in_code:
  p=d.add_paragraph(line); p.paragraph_format.space_after=Pt(2)
  for r in p.runs:r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
 elif line.startswith('# '): d.add_heading(line[2:],0)
 elif line.startswith('## '):
  p=d.add_heading(line[3:],1)
  p.paragraph_format.keep_with_next=True
 elif line.startswith('### '):d.add_heading(line[4:],2)
 elif line.startswith('- '):d.add_paragraph(line[2:],'List Bullet')
 elif line.strip():
  p=d.add_paragraph(line)
  if line.startswith('«'):
   for r in p.runs:r.italic=True
  if line.startswith('Archives disponibles'):
   for r in p.runs:r.font.size=Pt(9)
d.save(out.with_suffix('.docx'))
assert all(re.search(r'## '+f'{n:02d}'+r'\.',content) for n in range(1,20))
assert not re.search(r'\b[Jj](?:e |’)',content)
print(out, '19 modules ; 22 sous-tests ;',len(content.split()),'mots')
