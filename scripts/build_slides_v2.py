from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
R=Path(__file__).resolve().parents[1]; prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY='14283F'; TEAL='007F85'; DARK='23384A'; GREY='597082'; BG='F4F7FA'; AMBER='A76613'
def rect(s,x,y,w,h,color):
 q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=RGBColor.from_string(color); q.line.fill.background(); return q

def text(s,x,y,w,h,txt,size=22,color=DARK,bold=False):
 q=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=q.text_frame; tf.word_wrap=True
 tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='DejaVu Sans'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(10)
 return q

def slide(title,section='NOTRE ÉTUDE DVWA'):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 rect(s,0,0,13.333,.12,TEAL); text(s,.5,.3,12,.3,section,11,TEAL,True); text(s,.5,.88,12.3,.9,title,30,NAVY,True)
 rect(s,.5,7.03,12.3,.015,'CDDAE4'); text(s,.5,7.13,10,.2,'GROUPE 12  •  Rapport V2  •  Laboratoire local DVWA',10,GREY); text(s,12,7.1,.8,.3,f'{len(prs.slides):02d} / 28',10,GREY)
 return s

def notes(s,t): s.notes_slide.notes_text_frame.text=t

def bullets(s,items,x=.6,y=2,w=11.8,size=25):
 for i,t in enumerate(items): text(s,x,y+i*1.02,w,.91,'• '+t,size)

def pic(s,name,x,y,w,h,cap,crop=True,folder='attacks'):
 path=R/'docs/captures'/folder/name; iw,ih=Image.open(path).size
 # Cadrage natif PowerPoint : le fichier original reste embarqué intégralement.
 left=right=bottom=0
 if crop and folder=='attacks':
  if name.startswith(('T09-execution','T10-donnees','T17-','T19')): bottom=.80
  else:
   bottom=.40
   if not name.startswith(('T04-attaque','T05-','T06-attaque','T08-attaque','T15-')): left=.14; right=.14
 ratio=iw*(1-left-right)/(ih*(1-bottom)); dw=min(w,h*ratio); dh=dw/ratio
 q=s.shapes.add_picture(str(path), Inches(x+(w-dw)/2), Inches(y),width=Inches(dw),height=Inches(dh)); q.crop_left=left;q.crop_right=right;q.crop_bottom=bottom
 text(s,x,y+h+.08,w,.48,cap,12,GREY)
 return name
s=slide('Étude des attaques web et contre-mesures','PRÉSENTATION DU PROJET • GROUPE 12')
text(s,.65,2.15,7,1.6,'Notre laboratoire DVWA\nDe l’installation aux preuves',36,NAVY,True)
text(s,.65,4.35,7,1.1,'Deux instances séparées.\nDes résultats reliés aux captures.',25,TEAL)
text(s,.65,6.25,7,.4,'Version V2 • 15 septembre 2026',17,GREY)
rect(s,8.7,2,3.9,4.6,NAVY)
for y,a,b in [(2.4,'19','modules au périmètre'),(3.7,'22','sous-tests enregistrés'),(5,'39','captures dans le rapport')]:
 text(s,9,y,3.3,.6,a,34,'FFFFFF',True); text(s,9,y+.65,3.3,.45,b,16,'BFE7E7')
notes(s,'Nous présentons la synthèse du rapport V2. Les captures et les résultats d’attaques datent du 14 septembre ; la reprise du laboratoire date du 15 septembre. Source : Rapport-DVWA-V2.docx et evidence/attacks/results.json.')
s=slide('Nos objectifs et notre démarche')
bullets(s,['Nous disposons d’un environnement local reproductible.','Nous observons les effets des attaques sur la référence.','Nous relions chaque conclusion à une preuve.','Nous préparons les correctifs et leur validation par rejeu.'])
notes(s,'Plan : architecture et installation, méthode de preuve, résultats des modules, limites et prochaines étapes. Les correctifs personnels ne sont pas présentés comme déjà validés.')
s=slide('Notre architecture : cinq services isolés')
for x,y,w,h,c,title,sub in [(4.8,1.9,3.7,.8,NAVY,'Notre navigateur',''),(4.8,3.05,3.7,.9,TEAL,'gateway • Nginx','4280 / 4281 sur 127.0.0.1'),(.7,4.45,5.7,1.1,NAVY,'reference → reference-db','Réseau interne reference'),(6.9,4.45,5.7,1.1,NAVY,'corrections → corrections-db','Réseau interne corrections')]:
 rect(s,x,y,w,h,c); text(s,x+.15,y+.1,w-.3,.4,title,21,'FFFFFF',True)
 if sub: text(s,x+.15,y+.6,w-.3,.3,sub,14,'D4EBED')
text(s,6.3,2.75,.5,.3,'↓',21,TEAL);text(s,3.2,4,7,.4,'↙                                      ↘',25,TEAL)
text(s,.8,6,11.8,.7,'Nous publions uniquement le proxy en local ; les bases n’ont aucun port publié.',23)
notes(s,'Source : compose.yaml. Chaque base dispose de son volume. La copie corrections monte son répertoire vulnerabilities local en lecture seule.')
s=slide('Notre configuration conserve les données')
bullets(s,['Une base dvwa et un utilisateur dvwa par instance.','reference-db / corrections-db : persistance des bases.','reference-uploads / corrections-uploads : fichiers téléversés.','api-vendor : dépendances Composer de la copie de travail.'],size=24)
notes(s,'Source : compose.yaml et captures installation Setup Check. Les identifiants pédagogiques sont identiques, avec réseaux et volumes distincts. Nous n’avons pas réinitialisé les bases lors de la reprise.')
s=slide('Nous avons démarré et vérifié les deux instances')
text(s,.7,1.95,6.2,1.2,'5 services démarrés\n2 bases Healthy\n38 contrôles HTTP positifs',27,TEAL,True)
text(s,.7,4.25,6,1.5,'Référence : 127.0.0.1:4280\nCopie de travail : 127.0.0.1:4281\nConnexion : admin / password',21)
pic(s,'4280-accueil.png',7.2,1.95,5.4,4.4,'Accueil de la référence • capture du 14 septembre',False,'installation')
notes(s,'Reprise du 15 septembre : docker compose up -d puis docker compose ps et python3 scripts/check.py sans --init. 19 pages contrôlées par instance. Un contrôle HTTP positif confirme la disponibilité, pas une protection contre une attaque.')
s=slide('Notre méthode : de l’entrée à l’effet réel')
bullets(s,['Nous fixons l’instance, le niveau et le compte.','Nous comparons un usage normal et une entrée de test.','Nous conservons les captures et les résultats enregistrés.','Nous vérifions l’état final avant de conclure.'])
notes(s,'Sources : scripts/test_attacks.py, evidence/attacks/results.json et fichiers texte associés. Nous utilisons principalement low ; T15 et T18b utilisent medium. Les captures affichées sont parfois cadrées pour faciliter la lecture ; leurs originaux sont intégrés au PowerPoint.')
s=slide('Notre bilan : 22 sous-tests, trois types de résultats')
for x,n,l,c in [(.7,'17','Effets confirmés',TEAL),(4.95,'4','Essais à reprendre',AMBER),(9.2,'1','Persistance non confirmée',GREY)]:
 rect(s,x,2.1,3.5,2.4,c); text(s,x+.2,2.4,3.1,.9,n,56,'FFFFFF',True);text(s,x+.2,3.6,3.1,.7,l,21,'FFFFFF')
text(s,.8,5.2,11.7,1.1,'Nous distinguons l’exécution d’un test, l’effet observé et la validation d’un correctif.',27,NAVY,True)
notes(s,'22 entrées dans results.json. À reprendre : T06, T13, T18a, T19c. T19b est marqué success dans le journal, mais la relecture reste level=1 : nous ne comptons pas une élévation persistante. Aucun de ces nombres ne représente des correctifs personnels validés.')
data=[
('T01','Injection SQL','EFFET CONFIRMÉ','Nous passons de 1 à 5 utilisateurs affichés.','Une entrée modifie la sélection SQL.',['T01-normal.png','T01-attaque.png'],['Usage normal : un utilisateur','Injection : cinq utilisateurs']),
('T02','Injection SQL aveugle','EFFET CONFIRMÉ','Nous obtenons deux réponses selon la condition SQL.','Vrai : HTTP 200 ; faux : HTTP 404.',['T02-vrai.png','T02-faux.png'],['Condition vraie : utilisateur présent','Condition fausse : utilisateur absent']),
('T03','Injection de commandes','EFFET CONFIRMÉ','Nous retrouvons G12_COMMAND_OK après le ping.','Notre entrée est interprétée par le shell.',['T03-normal.png','T03-attaque.png'],['Ping normal','Commande ajoutée : marqueur visible']),
('T04','XSS réfléchie','EFFET CONFIRMÉ','Nous observons le marqueur ajouté par JavaScript.','La charge provient de notre requête.',['T04-normal.png','T04-attaque.png'],['Affichage du nom saisi','Marqueur G12_XSS_R en haut de page']),
('T05','XSS stockée','EFFET CONFIRMÉ','Nous retrouvons l’exécution dans une seconde session.','Aucune nouvelle publication n’est nécessaire.',['T05-publication.png','T05-visiteur-distinct.png'],['Publication de la charge','Visite dans une session distincte']),
('T06','XSS DOM','À REPRENDRE','Nous conservons les captures, mais le contrôle a échoué.','Nous ne confirmons pas l’exécution JavaScript.',['T06-normal.png','T06-attaque.png'],['Sélection normale de la langue','État obtenu après la charge']),
('T07','Requête intersite CSRF','EFFET CONFIRMÉ','Nous changeons le mot de passe depuis le port 4290.','La reconnexion confirme le changement.',['T07-changement.png','T07-reconnexion.png'],['Message de changement du mot de passe','Reconnexion avec le nouveau mot de passe']),
('T08','Inclusion de fichiers','EFFET CONFIRMÉ','Nous obtenons robots.txt par un chemin relatif.','Le fichier se situe hors du dossier du module.',['T08-normal.png','T08-attaque.png'],['Page incluse normalement','Contenu de robots.txt et avertissements']),
('T09','Téléversement de fichier PHP','EFFET CONFIRMÉ','Nous téléversons puis appelons un fichier témoin PHP.','La réponse contient G12_UPLOAD_EXECUTED.',['T09-televersement.png','T09-execution.png'],['Téléversement accepté','Exécution du fichier : marqueur retourné']),
('T10','Contournement d’autorisation','EFFET CONFIRMÉ','Nous utilisons le compte ordinaire gordonb.','La route de données expose cinq utilisateurs.',['T10-compte-ordinaire.png','T10-donnees.png'],['Session du compte ordinaire','Données obtenues directement']),
('T11','Contrôle d’accès et cookie','EFFET CONFIRMÉ','Nous remplaçons le cookie user_id=2 par user_id=1.','Le profil initialement refusé devient accessible.',['T11-refus.png','T11-cookie-falsifie.png'],['Accès refusé avec le cookie initial','Profil 1 affiché après falsification']),
('T12','Essais de mots de passe','EFFET CONFIRMÉ','Nous obtenons deux échecs puis un succès.','Notre conclusion porte sur trois essais seulement.',['T12-essai-1.png','T12-essai-3.png'],['Premier mot de passe incorrect','Mot de passe connu accepté']),
('T13','Identifiants de session faibles','À REPRENDRE','Nous avons ouvert le module.','Une navigation a interrompu la collecte des jetons.',['T13-module.png'],['Module accessible ; prédictibilité non validée']),
('T14','Parcours CAPTCHA','EFFET CONFIRMÉ','Nous soumettons directement step=2.','La reconnexion confirme le nouveau mot de passe.',['T14-saut-etape.png','T14-reconnexion.png'],['État après soumission directe de step=2','Reconnexion enregistrée après changement']),
('T15','Réutilisation d’un nonce CSP','EFFET CONFIRMÉ • MEDIUM','Nous réutilisons le nonce statique dans un script.','Le marqueur G12_CSP_EXECUTED apparaît.',['T15-nonce-reutilise.png'],['Script accepté avec le nonce connu']),
('T16','Jeton calculé côté client','EFFET CONFIRMÉ','Nous recalculons le jeton puis soumettons la phrase.','Le module affiche « Well done! ».',['T16-jeton-initial.png','T16-jeton-recalcule.png'],['Jeton initial : échec','Jeton recalculé : succès']),
('T17','Redirection ouverte','EFFET CONFIRMÉ','Nous fournissons une destination locale contrôlée.','HTTP 302 dirige le navigateur vers le port 4290.',['T17-destination.png'],['Page témoin atteinte par le navigateur']),
('T18','Cryptographie : deux sous-tests','RÉSULTATS DISTINCTS','XOR : nous conservons le test à reprendre.','ECB : le jeton recomposé est accepté comme administrateur.',['T18a-xor-decode.png','T18b-blocs-recombines.png'],['T18a • Décodage visible, validation incomplète','T18b • « Welcome administrator Sweep »']),
('T19','API : nous vérifions les réponses et l’état','RÉSULTATS DISTINCTS','v1 expose les hachages password absents de v2.','PUT : level=0 ; relecture : level=1. Persistance non confirmée.',['T19a-v1.png','T19b-relecture.png'],['T19a • Données sensibles dans v1','T19b • Relecture : level reste à 1'])]
entries=json.loads((R/'evidence/attacks/results.json').read_text())
for ident,title,status,claim,limit,files,caps in data:
 s=slide(f'{ident} — {title}','NOS EXPÉRIMENTATIONS • RÉFÉRENCE LOCALE')
 color=TEAL if status.startswith('EFFET') else AMBER
 text(s,.6,1.7,12,.35,status,13,color,True)
 text(s,.6,2.13,12.1,.57,claim,24,NAVY,True); text(s,.6,2.8,12.1,.64,limit,20)
 if len(files)==2:
  for i,(f,c) in enumerate(zip(files,caps)): pic(s,f,.6+i*6.35,3.55,5.8,2.72,c)
 else: pic(s,files[0],2.6,3.55,8.1,2.72,caps[0])
 matches=[e for e in entries if e['id'].startswith(ident)]
 extra={'T07':'Les origines diffèrent par le port mais restent same-site. Nous ne généralisons pas à tous les contextes intersites.','T12':'Le dernier essai utilise le mot de passe connu. Trois essais ne prouvent pas l’absence générale de limitation.','T14':'Aucun CAPTCHA externe n’a été résolu. Le parcours normal avec reCAPTCHA reste à valider.','T16':'Ce test démontre la contrôlabilité du client ; il ne démontre pas une XSS distante.','T18':'T18a : le journal indique un échec de vérification, malgré le message visible. T18b utilise medium. Aucun oracle complet ni rupture de l’algorithme AES n’est démontré.','T19':'T19c : essai d’injection de commandes API à reprendre, sans capture ni effet confirmé. Le succès enregistré pour T19b ne confirme pas une élévation persistante.'}.get(ident,'Conclusion limitée au scénario enregistré. Aucun rejeu après correctif personnel n’est documenté ici.')
 notes(s,extra+'\nSources : evidence/attacks/results.json ; scripts/test_attacks.py. Captures : '+', '.join('docs/captures/attacks/'+f for f in files)+'\nRésultats conservés : '+json.dumps(matches,ensure_ascii=False))
s=slide('Nos prochaines étapes : corriger puis rejouer')
bullets(s,['Nous reprenons T06, T13, T18a et T19c.','Nous vérifions la mutation et la relecture de T19b.','Nous identifions chaque correctif personnel dans Git.','Nous rejouons l’attaque et l’usage légitime après correction.'],size=24)
notes(s,'Ces étapes sont futures. Les protections pédagogiques de DVWA restent attribuées à ses auteurs. Nous ne présentons pas une sélection du niveau impossible comme notre contribution ni comme une sécurisation complète des endpoints.')
s=slide('Notre conclusion')
text(s,.7,2,11.8,1.2,'Nous disposons d’un laboratoire opérationnel\net de preuves pour analyser les effets observés.',31,NAVY,True)
bullets(s,['Nous séparons disponibilité, exploitation et correction.','Nous signalons les résultats incomplets.','Nous fondons la suite du travail sur le rejeu et l’état réel.'],y=3.6,size=24)
text(s,.8,6.7,11.7,.25,'Sources : rapport V2 • compose.yaml • evidence/attacks/results.json • captures du projet',11,GREY)
notes(s,'Référence : GROUPE-12/Rapport-DVWA-V2.docx et .pdf. Projet DVWA : https://github.com/digininja/DVWA ; commit des sources locales b496a5d3de6b967410155e1b7d3e51e9d035eb22. Nous avons utilisé les preuves locales disponibles, sans rejouer les attaques pour produire cette présentation. Merci pour votre attention.')
assert len(prs.slides)==28
out=R/'GROUPE-12/Presentation-DVWA-V2.pptx'; prs.save(out); print(out)
