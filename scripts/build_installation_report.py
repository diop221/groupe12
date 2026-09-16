from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
R=Path(__file__).resolve().parents[1]
D=Document(); sec=D.sections[0]
sec.page_height=Cm(29.7); sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(1.8); sec.left_margin=sec.right_margin=Cm(2)
normal=D.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
 D.styles[name].font.color.rgb=RGBColor.from_string('17365D')
for name in ['Code','FigureCaption','TableCaption']:
 s=D.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Consolas' if name=='Code' else 'Calibri'; s.font.size=Pt(8 if name=='Code' else 10)
 s.paragraph_format.space_after=Pt(4)
 if name!='Code': s.font.italic=True
D.styles['Code'].paragraph_format.left_indent=Cm(.25)
header=sec.header.paragraphs[0]; header.text='GROUPE 12  •  Étude des attaques web  •  Laboratoire DVWA'; header.style='Caption'
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Groupe 12 — Page ')
def fld(p,code,fallback=''):
 r=p.add_run(); b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin'); r._r.append(b)
 t=OxmlElement('w:instrText'); t.set(qn('xml:space'),'preserve'); t.text=' '+code+' '; r._r.append(t)
 x=OxmlElement('w:fldChar'); x.set(qn('w:fldCharType'),'separate'); r._r.append(x)
 p.add_run(fallback)
 r=p.add_run(); e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end'); r._r.append(e)
fld(f,'PAGE','1')
def p(t): D.add_paragraph(t)
def h(t,level=1): D.add_heading(t,level)
def code(t):
 for line in t.strip().splitlines(): D.add_paragraph(line,'Code')
def page(): D.add_page_break()
figs=[]; tabs=[]
def figure(file,caption,description):
 path=R/'docs/captures/installation'/file
 # Fit full-page browser captures within a page, without stretching.
 import struct
 raw=path.read_bytes(); w,ht=struct.unpack('>II',raw[16:24]); width=min(6.65,7.2*w/ht)
 pp=D.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.keep_with_next=True
 pp.add_run().add_picture(str(path),width=Inches(width))
 figs.append(caption); D.add_paragraph(f'Figure {len(figs)} — {caption}','FigureCaption'); p(description)
def table(title,heads,rows):
 tabs.append(title); cp=D.add_paragraph(f'Tableau {len(tabs)} — {title}','TableCaption'); cp.paragraph_format.keep_with_next=True
 t=D.add_table(rows=1,cols=len(heads)); t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,heads): c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row): c.text=str(v)
 p('')
D.add_paragraph('GROUPE 12','Subtitle')
D.add_heading('Étude et implémentation des attaques web et contre-mesures',0)
D.add_paragraph('Installation et configuration du laboratoire DVWA\nDe la préparation de l’environnement à l’affichage de l’interface','Subtitle')
p('Rapport technique — 14 septembre 2026')
p('Ce document décrit la mise en place effectivement retrouvée dans le projet, les commandes permettant de la reproduire et la vérification actuelle de l’accès aux deux interfaces. Les captures ont été prises sur les instances locales le 14 septembre 2026.')
p('Périmètre de cette version : installation, configuration, initialisation, connexion et contrôles de disponibilité. L’expérimentation des attaques et la validation de correctifs feront l’objet de la suite du rapport.')
page()
h('Sommaire'); fld(D.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u','Sommaire à actualiser.')
page(); h('Liste des figures'); fld(D.add_paragraph(),'TOC \\t "FigureCaption,1" \\h \\z','Liste à actualiser.')
h('Liste des tableaux'); fld(D.add_paragraph(),'TOC \\t "TableCaption,1" \\h \\z','Liste à actualiser.')
page(); h('1. Présentation du projet')
p('Le projet porte sur l’étude pratique des attaques web, notamment les injections SQL, et de leurs contre-mesures. Pour observer ces mécanismes dans un environnement maîtrisé, nous utilisons Damn Vulnerable Web Application (DVWA), une application PHP avec une base MySQL/MariaDB conçue pour l’apprentissage de la sécurité web.')
p('Notre première étape consiste à disposer d’une plateforme accessible, reproductible et séparée des autres projets. Le déploiement retenu repose sur Docker Compose : les composants web, les bases et le proxy sont exécutés dans des conteneurs. Il n’est donc pas nécessaire d’installer Apache, PHP et MariaDB directement sur Ubuntu pour ce laboratoire.')
h('1.1. Objectifs de la mise en place',2)
p('La mise en place doit permettre de se connecter à DVWA, d’afficher les modules pédagogiques, de conserver les données entre deux redémarrages et de préparer une comparaison entre une référence et une copie destinée aux corrections. Les deux sites sont initialement vulnérables et utilisent le niveau low.')
h('1.2. Origine des informations et traçabilité',2)
p('Le laboratoire initial a été préparé le 12 septembre 2026. Le présent rapport s’appuie sur README.md, compose.yaml, les fichiers infra/, le dépôt Git, les scripts et les preuves du dossier evidence/. Le démarrage, la connexion et les contrôles ont été rejoués le 14 septembre. Les bases existantes ont été conservées.')
p('L’historique complet de la première installation n’est pas disponible. Les commandes Git de création et de préparation sont donc une procédure de reproduction conforme à l’état constaté, et non une transcription garantie de chaque commande saisie le 12 septembre. Les commandes de démarrage et de vérification présentées ont été exécutées lors de la reprise.')
h('2. Environnement et architecture')
table('Environnement relevé le 14 septembre 2026',['Composant','Version / rôle'],[['Système hôte','Ubuntu 26.04 LTS'],['Docker','29.1.3, paquet Ubuntu docker.io'],['Docker Compose','2.40.3, paquet docker-compose-v2'],['Git','2.55.0'],['Python hôte','3.14.4 — script de contrôle HTTP'],['PHP dans DVWA','8.5.10'],['Client MariaDB du conteneur','11.4.13'],['Sources locales DVWA','b496a5d3de6b967410155e1b7d3e51e9d035eb22']])
p('Les versions du moteur, de Compose et des outils ont été relevées avec les commandes suivantes :')
code('cat /etc/os-release\ndocker --version\ndocker compose version\ngit --version\npython3 --version\ndpkg-query -W docker.io docker-compose-v2')
h('2.1. Rôle des cinq services',2)
table('Services et séparation des instances',['Service','Fonction','Accès / réseau'],[['gateway','Proxy Nginx','127.0.0.1:4280 et :4281'],['reference','DVWA de référence','Réseau reference ; HTTP interne 80'],['reference-db','Base de la référence','Réseau reference ; 3306 interne'],['corrections','Copie de travail DVWA','Réseau corrections ; HTTP interne 80'],['corrections-db','Base de la copie de travail','Réseau corrections ; 3306 interne']])
p('Le navigateur contacte Nginx sur la boucle locale. Le port 4280 est transmis au service reference, et le port 4281 au service corrections. Les réseaux reference et corrections sont déclarés internes. Les bases ne possèdent aucun port publié sur la machine hôte. Les données et les téléversements sont conservés dans des volumes distincts.')
p('La référence web provient d’une image identifiée par son digest. La copie de travail utilise la même image, mais son dossier vulnerabilities est remplacé par le répertoire local corrections/vulnerabilities. Le commit identifie les sources locales ; le digest identifie séparément l’image déployée.')
h('3. Préparation des fichiers du laboratoire')
h('3.1. Vérification des prérequis',2)
p('Docker et son module Compose sont déjà présents sur la machine examinée. Le relevé des paquets confirme docker.io et docker-compose-v2. Nous ne présentons donc pas une réinstallation de Docker comme une étape exécutée pendant cette reprise. Avant de démarrer le laboratoire, les commandes de version du chapitre précédent permettent de vérifier leur disponibilité.')
h('3.2. Récupération et identification des sources',2)
p('Pour reproduire l’organisation sur une installation neuve, créer le dossier du projet, récupérer le dépôt officiel et se placer sur le commit étudié. Ces commandes de création ne sont pas à rejouer sur les dossiers déjà existants :')
code('mkdir -p /home/modou/dvwa-study\ncd /home/modou/dvwa-study\ngit clone https://github.com/digininja/DVWA.git upstream\ngit -C upstream checkout b496a5d3de6b967410155e1b7d3e51e9d035eb22\ngit -C upstream worktree add -b study/corrections ../corrections b496a5d3de6b967410155e1b7d3e51e9d035eb22')
p('Le worktree crée un second répertoire de travail associé au même dépôt Git, sur la branche study/corrections. Il permet de conserver la référence tout en préparant des modifications isolées. Au moment du contrôle, les deux copies pointent sur le même commit et aucune modification personnelle n’est présente dans corrections.')
code('git -C upstream worktree list\ngit -C corrections status --short\ngit -C corrections log -1 --oneline')
h('3.3. Organisation du projet',2)
code('dvwa-study/\n  compose.yaml       # orchestration des cinq services\n  VERSION            # commit des sources locales\n  upstream/          # sources de référence\n  corrections/       # copie de travail Git\n  infra/             # réglages PHP et Nginx\n  scripts/check.py   # initialisation et contrôles HTTP\n  evidence/          # résultats et journal de démarrage\n  docs/              # inventaire, captures et script vidéo\n  GROUPE-12/         # fichiers de remise')
p('Les fichiers de configuration complets sont reproduits en annexe. Dans un nouveau dossier, ils doivent être enregistrés aux chemins indiqués avant l’exécution de Docker Compose. Le répertoire utilisé par le module API doit également exister :')
code('cd /home/modou/dvwa-study\nmkdir -p infra evidence docs\nmkdir -p corrections/vulnerabilities/api/vendor')
h('4. Configuration des services')
h('4.1. Configuration des bases MariaDB',2)
p('Chaque instance dispose d’une base dédiée appelée dvwa et d’un utilisateur applicatif dvwa. Les paramètres de démonstration sont déclarés dans compose.yaml. Les deux bases utilisent les mêmes identifiants pédagogiques mais des volumes et réseaux différents.')
code('MARIADB_ROOT_PASSWORD: local-study-root\nMARIADB_DATABASE: dvwa\nMARIADB_USER: dvwa\nMARIADB_PASSWORD: p@ssw0rd')
p('DVWA reçoit DB_SERVER=reference-db ou DB_SERVER=corrections-db. Les autres valeurs par défaut du fichier config/config.inc.php.dist correspondent à la base dvwa, à l’utilisateur dvwa, au mot de passe p@ssw0rd et au port 3306. Le healthcheck MariaDB vérifie la connexion et l’initialisation d’InnoDB avant le démarrage du service web dépendant.')
h('4.2. Configuration de PHP et des sessions',2)
p('Le fichier infra/lab.ini active les fonctionnalités nécessaires à certains exercices. Il est monté en lecture seule dans /usr/local/etc/php/conf.d/zz-lab.ini :')
code((R/'infra/lab.ini').read_text())
p('allow_url_fopen et allow_url_include servent aux exercices d’inclusion ; display_errors rend les erreurs visibles dans ce contexte pédagogique. La taille maximale d’un téléversement est de 8 Mo et celle d’une requête POST de 10 Mo. PHP 8.5.10 signale que allow_url_include est obsolète : cet avertissement a été observé et n’empêche pas l’affichage de l’interface.')
p('Les fichiers reference.ini et corrections.ini attribuent des noms de session distincts :')
code('session.name=DVWA_REFERENCE\n# Dans le fichier de la seconde instance :\nsession.name=DVWA_CORRECTIONS')
p('Le cookie security peut néanmoins être partagé entre ports d’une même adresse. Pour les futures comparaisons de niveaux, ouvrir les deux instances dans des profils de navigateur distincts. DEFAULT_SECURITY_LEVEL=low fixe le niveau initial des nouvelles sessions.')
h('4.3. Configuration de Nginx et des ports',2)
p('Le fichier infra/nginx.conf contient deux serveurs : le premier écoute sur 4280 et transmet les requêtes à reference:80 ; le second écoute sur 4281 et les transmet à corrections:80. La directive client_max_body_size 10m correspond à la limite de requête PHP. Le fichier complet apparaît en annexe.')
code("ports:\n  - '127.0.0.1:4280:4280'\n  - '127.0.0.1:4281:4281'")
p('La liaison à 127.0.0.1 rend ces adresses accessibles depuis la machine du laboratoire. L’adresse à saisir dans le navigateur est celle du proxy, et non le nom interne du conteneur.')
h('4.4. Volumes et dépendances du module API',2)
p('Les volumes reference-db et corrections-db conservent les bases. Les volumes reference-uploads et corrections-uploads conservent les fichiers téléversés. Le volume api-vendor fournit un emplacement distinct aux dépendances Composer de la copie de travail. La page Setup Check confirme actuellement que les dépendances API sont installées et que mod_rewrite est activé.')
h('5. Démarrage du laboratoire')
p('Depuis le dossier du projet, lancer les services en arrière-plan :')
code('cd /home/modou/dvwa-study\ndocker compose up -d')
p('Docker Compose lit compose.yaml, crée ou réutilise les réseaux et volumes, puis démarre les bases. Lorsque leurs contrôles de santé deviennent positifs, les services web démarrent, suivis du proxy. Si les images ne sont pas présentes, Docker les télécharge. Le journal initial evidence/startup.log conserve des traces de cette récupération.')
p('Lors de la reprise du 14 septembre, la commande a recréé les cinq conteneurs et les a démarrés. Les deux bases sont passées à l’état Healthy. Les volumes existants ont été conservés ; aucune commande de réinitialisation des bases n’a été exécutée.')
code('docker compose ps\ndocker compose exec -T reference php -v\ndocker compose exec -T reference-db mariadb --version')
table('Résultat du contrôle de démarrage du 14 septembre',['Service','État observé'],[['reference-db','Up — healthy'],['corrections-db','Up — healthy'],['reference','Up'],['corrections','Up'],['gateway','Up — ports locaux 4280 et 4281 publiés']])
p('Si Docker refuse l’accès à son socket, la commande doit être exécutée dans un contexte disposant des droits Docker. Dans cette session, le contrôle et le démarrage ont nécessité une autorisation d’accès au moteur. Une erreur de permission n’indique pas que la configuration DVWA est incorrecte.')
h('6. Initialisation et contrôle de configuration')
h('6.1. Initialisation des données de démonstration',2)
p('Lors de la première installation, accéder à http://127.0.0.1:4280/setup.php, vérifier les paramètres puis cliquer sur Create / Reset Database. Répéter pour http://127.0.0.1:4281/setup.php. Cette action crée ou réinitialise les tables et les comptes de démonstration. Elle ne doit pas être répétée simplement pour prendre une capture lorsque la base contient déjà des travaux à conserver.')
p('Le projet fournit également la commande ci-dessous. Le script charge setup.php, récupère le jeton du formulaire et soumet l’action officielle sur les deux instances avant de vérifier leur connexion :')
code('python3 scripts/check.py --init')
p('Cette commande est la procédure d’initialisation documentée du projet. Elle n’a pas été rejouée le 14 septembre : les connexions fonctionnaient déjà avec les données conservées. La capture suivante montre le contrôle de configuration actuel, et non une nouvelle réinitialisation.')
page(); figure('4280-configuration.png','Contrôle de configuration de la référence, port 4280','La page Database Setup affiche PHP 8.5.10, les extensions mysqli et pdo_mysql, mod_rewrite activé, le serveur reference-db et les dépendances API installées. La clé reCAPTCHA apparaît Missing. La mention DVWA version: Unknown de la page ne remplace pas l’identification des sources par Git et de l’image par digest.')
page(); figure('4281-configuration.png','Contrôle de configuration de la copie de travail, port 4281','Le serveur de base affiché est corrections-db : la seconde application utilise sa propre base. Les clés reCAPTCHA sont également absentes. L’absence de ces clés limite le parcours CAPTCHA et ne bloque pas la connexion générale.')
h('7. Connexion et apparition de l’interface')
h('7.1. Ouverture du formulaire de connexion',2)
p('Sur la machine du laboratoire, ouvrir le navigateur et saisir http://127.0.0.1:4280/login.php. Le formulaire DVWA présente les champs Username et Password ainsi que le bouton Login. Pour la copie de travail, utiliser la même adresse avec le port 4281.')
table('Adresses et compte de démonstration',['Élément','Valeur'],[['Référence','http://127.0.0.1:4280'],['Copie de travail','http://127.0.0.1:4281'],['Nom d’utilisateur initial','admin'],['Mot de passe initial','password']])
figure('4280-connexion.png','Formulaire de connexion DVWA sur la référence','Capture réelle avant saisie des identifiants. L’affichage du formulaire confirme que le navigateur atteint le service web par le proxy ; il ne prouve pas encore une connexion à la base réussie.')
h('7.2. Authentification et page d’accueil',2)
p('Saisir admin dans Username et password dans Password, puis cliquer sur Login. Après validation, DVWA affiche sa page d’accueil. Lors du contrôle, cette opération a réussi sur les deux ports avec des contextes de navigateur séparés.')
page(); figure('4280-accueil.png','Interface d’accueil après connexion sur la référence','La page Welcome to Damn Vulnerable Web Application!, le menu latéral, le message de connexion en tant que admin et le niveau low sont visibles. Cette capture matérialise l’objectif de la mise en place : l’interface authentifiée est opérationnelle.')
page(); figure('4281-accueil.png','Interface d’accueil après connexion sur la copie de travail','La seconde instance affiche également la page d’accueil avec le compte admin. Elle constitue le support des futures modifications ; elle reste vulnérable à cette étape.')
h('7.3. Vérification du niveau de sécurité',2)
p('Cliquer sur DVWA Security dans le menu latéral. La page indique le niveau courant et propose low, medium, high et impossible. Le niveau courant observé est low, conformément à la configuration initiale. Les protections des autres niveaux sont fournies par DVWA et ne constituent pas des corrections réalisées par notre groupe.')
figure('4280-securite.png','Page DVWA Security et niveau courant low','Le niveau visible permet de fixer les conditions des futurs essais. Aucun changement de niveau n’a été nécessaire pour les captures de cette installation.')
h('7.4. Ouverture d’un module pédagogique',2)
p('Cliquer sur SQL Injection dans le menu. Le navigateur ouvre /vulnerabilities/sqli/ et affiche le champ User ID ainsi que le bouton Submit. Cette étape vérifie la navigation vers un exercice ; aucune charge d’injection n’a été soumise dans ce contrôle d’installation.')
figure('4280-module-sqli.png','Affichage du module SQL Injection avant toute attaque','Le formulaire de l’exercice est accessible après authentification. La capture prouve l’ouverture du module et ne doit pas être interprétée comme un résultat d’exploitation.')
h('8. Vérification de disponibilité et bilan')
h('8.1. Contrôle automatisé des pages',2)
p('La commande suivante réalise un contrôle sans réinitialiser les bases :')
code('cd /home/modou/dvwa-study\npython3 scripts/check.py')
p('Le script crée une session HTTP distincte par instance, récupère le jeton de connexion, soumet admin/password puis parcourt les 19 répertoires de modules. Pour l’inclusion de fichiers, il utilise le paramètre page=include.php. Il vérifie un statut HTTP 200, l’absence de redirection vers login.php et l’absence du texte Fatal error. Les résultats sont enregistrés dans evidence/http-check.json.')
results=json.loads((R/'evidence/http-check.json').read_text())
table('Contrôle HTTP rejoué le 14 septembre 2026',['Instance','Pages contrôlées','Résultats positifs'],[[str(port),str(sum(x['port']==port for x in results)),str(sum(x['port']==port and x['page_ok'] for x in results))] for port in (4280,4281)])
p('Les 38 contrôles sont positifs. Le contrôle navigateur complète ce résultat en montrant les pages rendues et la connexion effective. Ces vérifications ne couvrent pas l’exploitation des failles, l’exécution de JavaScript dans tous les modules ni l’efficacité de contre-mesures.')
h('8.2. Arrêt, reprise et diagnostic',2)
code('docker compose down\n# Reprise ultérieure :\ndocker compose up -d\ndocker compose ps\n# Diagnostic en cas de problème :\ndocker compose logs --tail=100 reference reference-db gateway')
p('docker compose down arrête et supprime les conteneurs et réseaux du projet, tout en conservant les volumes nommés par défaut. Ne pas ajouter -v si l’objectif est de préserver les bases et fichiers téléversés. Les commandes d’arrêt ci-dessus sont fournies pour l’exploitation ultérieure ; le laboratoire a été laissé démarré à l’issue de la vérification.')
h('8.3. Résultat de la mise en place',2)
p('Le laboratoire est opérationnel jusqu’à l’affichage de l’interface : les cinq services sont démarrés, les bases sont disponibles, la connexion admin fonctionne sur les deux instances et les pages de modules sont accessibles. Les fichiers de configuration et les commandes de reprise sont conservés dans le projet.')
p('La suite du travail consistera à décrire et réaliser chaque attaque retenue, enregistrer des captures avant et après correction et tester les usages légitimes. Aucune conclusion d’efficacité des protections personnelles n’est formulée dans ce rapport d’installation.')
page(); h('Annexe A. Fichier compose.yaml complet')
p('Enregistrer ce contenu dans /home/modou/dvwa-study/compose.yaml. Les digests ci-dessous sont ceux du fichier actuellement utilisé ; ils fixent les images du laboratoire.')
code((R/'compose.yaml').read_text())
h('Annexe B. Fichiers de configuration complémentaires')
for filename in ['nginx.conf','lab.ini','reference.ini','corrections.ini']:
 h('infra/'+filename,2); code((R/'infra'/filename).read_text())
h('Annexe C. Preuves et sources du rapport')
p('Les captures originales sont conservées dans docs/captures/installation/. manifest.json associe chaque image à son URL et au titre de la page. Les images du rapport sont des captures du contenu de pages réellement visitées, sans reconstitution graphique de l’interface.')
p('Éléments locaux utilisés : README.md ; VERSION ; compose.yaml ; infra/ ; upstream/config/config.inc.php.dist ; scripts/check.py ; evidence/startup.log ; evidence/http-check.json ; état Git des deux worktrees.')
p('Dépôt d’origine : https://github.com/digininja/DVWA. Le rapport décrit la copie et les fichiers locaux examinés ; il ne dépend pas d’une vérification du contenu actuel du site distant.')
# avoid including the front matter in the main TOC
for para in D.paragraphs:
 if para.text in ['Sommaire','Liste des figures','Liste des tableaux']:
  para.style='Title'
settings=D.settings.element
u=OxmlElement('w:updateFields'); u.set(qn('w:val'),'true'); settings.append(u)
out=R/'GROUPE-12/Rapport-DVWA.docx'; D.save(out)
print(out, 'Figures:',len(figs),'Tableaux:',len(tabs))
