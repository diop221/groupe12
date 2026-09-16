"""Crée une V2 indépendante à partir du rapport et des preuves existantes."""
from pathlib import Path
import json, re, struct
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[1]
d=Document(R/'GROUPE-12/Rapport-DVWA.docx')
replacements={
'Installation et configuration du laboratoire DVWA':'Notre laboratoire DVWA : installation et expérimentation\nRapport du groupe 12 — Version V2',
'Rapport technique —':'Rapport technique — Version V2 — 15 septembre 2026',
'Ce document décrit':'Dans ce rapport, nous présentons notre laboratoire DVWA, sa configuration et les essais documentés dans notre projet. Nous intégrons les captures d’installation et les 39 captures du dossier docs/captures/attacks, accompagnées des résultats enregistrés le 14 septembre 2026.',
'Périmètre de cette version':'Dans cette V2, nous couvrons la mise en place, la connexion, la disponibilité et les expérimentations conservées. Nous distinguons les effets confirmés, les essais à reprendre et les correctifs dont la validation reste à réaliser.',
'Le laboratoire initial a été préparé':'Nous nous appuyons sur les fichiers du laboratoire préparé le 12 septembre 2026, sur les contrôles et captures du 14 septembre, et sur le redémarrage du 15 septembre. Nous conservons les preuves dans evidence/ et docs/captures/.',
'L’historique complet':'Nous ne disposons pas de l’historique complet de la première installation. Nous présentons donc les commandes de création comme une procédure de reproduction ; nos résultats de reprise et d’expérimentation sont rattachés aux journaux disponibles.',
'Docker et son module Compose':'Nous utilisons Docker et Docker Compose déjà installés sur notre machine. Nous vérifions leur présence avec les commandes de version ci-dessus avant de démarrer le laboratoire.',
'Pour reproduire l’organisation':'Pour reproduire notre organisation sur une installation neuve, nous créons le dossier du projet, récupérons les sources officielles et sélectionnons le commit étudié. Nous réservons ces commandes à un nouveau dossier :',
'Chaque instance dispose':'Nous attribuons à chaque instance une base dédiée nommée dvwa et un utilisateur applicatif dvwa. Nous déclarons les paramètres dans compose.yaml. Nous utilisons les mêmes identifiants pédagogiques avec des volumes et des réseaux distincts.',
'Les volumes reference-db':'Nous conservons les bases dans reference-db et corrections-db, et les fichiers téléversés dans reference-uploads et corrections-uploads. Nous réservons api-vendor aux dépendances Composer de notre copie de travail. Sur les captures Setup Check, nous constatons que les dépendances API sont installées et que mod_rewrite est activé.',
'Depuis le dossier du projet':'Depuis le dossier de notre projet, nous démarrons les services en arrière-plan avec la commande suivante :',
'Lors de la reprise du 14':'Lors de notre reprise du 15 septembre, Docker Compose a recréé les deux conteneurs web et le proxy, puis démarré les cinq services. Nous avons vérifié que les deux bases étaient Healthy. Nous avons conservé les volumes existants sans réinitialiser les données.',
'Si Docker refuse':'Pour exécuter nos commandes Docker, nous utilisons un compte disposant des droits d’accès au moteur. Un refus d’accès au socket doit être résolu avant de vérifier les services.',
'Lors de la première installation, accéder':'Pour initialiser une installation neuve, nous ouvrons setup.php sur chaque instance et utilisons Create / Reset Database. Cette action crée ou réinitialise les tables et comptes de démonstration. Nous ne la rejouons pas lors d’une simple reprise avec des données à conserver.',
'Le projet fournit également':'Nous pouvons également utiliser la commande ci-dessous : notre script charge setup.php, récupère le jeton du formulaire et soumet l’action officielle sur les deux instances avant de vérifier leur connexion.',
'Cette commande est la procédure':'Nous réservons --init à une initialisation volontaire. Nous ne l’avons pas utilisé lors de la reprise du 15 septembre, car nos connexions fonctionnaient avec les données conservées. Les captures suivantes montrent les contrôles de configuration du 14 septembre.',
'Sur la machine du laboratoire, ouvrir':'Sur notre machine, nous ouvrons http://127.0.0.1:4280/login.php pour la référence et http://127.0.0.1:4281/login.php pour la copie de travail. Nous retrouvons les champs Username et Password et le bouton Login.',
'Saisir admin dans':'Nous saisissons admin dans Username et password dans Password, puis nous cliquons sur Login. Nous obtenons la page d’accueil sur les deux instances avec des contextes de navigateur séparés.',
'Cliquer sur DVWA Security':'Nous ouvrons DVWA Security pour relever le niveau courant. Les captures d’installation montrent low. Les autres niveaux sont des protections pédagogiques fournies par DVWA ; nous les distinguons de nos propres contributions.',
'Cliquer sur SQL Injection':'Nous ouvrons SQL Injection dans le menu. Nous retrouvons le champ User ID et le bouton Submit. Cette capture documente notre accès initial au module ; nous présentons les essais d’injection au chapitre 9.',
'La commande suivante réalise':'Nous contrôlons les pages sans réinitialiser nos bases avec la commande suivante :',
'Les 38 contrôles sont positifs':'Nous avons obtenu 38 contrôles positifs lors de la reprise du 15 septembre : 19 pages par instance. Nous utilisons ce résultat pour confirmer la disponibilité ; nos conclusions sur les attaques reposent sur les preuves distinctes du chapitre 9.',
'Le laboratoire est opérationnel':'Nous disposons d’un laboratoire opérationnel : nos cinq services sont démarrés, nos bases sont disponibles et nous pouvons nous connecter aux deux instances. Nous conservons nos configurations et commandes de reprise dans le projet.',
'La suite du travail consistera':'Nous complétons maintenant ce bilan d’installation par nos expérimentations illustrées au chapitre 9. Nous réservons la validation de correctifs personnels à des modifications identifiées et à des tests de rejeu documentés.',
'Enregistrer ce contenu':'Nous conservons cette configuration dans /home/modou/groupe12/compose.yaml. Les digests fixent les images utilisées dans notre laboratoire.',
'Les captures originales sont conservées':'Nous conservons les captures d’installation dans docs/captures/installation/ et les 39 captures d’essais dans docs/captures/attacks/. Nous les intégrons sans modifier leur contenu. Pour les essais, evidence/attacks/results.json associe les images aux URL, aux dates et aux résultats ; les fichiers texte voisins conservent le texte des pages.',
'Éléments locaux utilisés':'Nous avons utilisé README.md, VERSION, compose.yaml, infra/, scripts/check.py, scripts/test_attacks.py, evidence/http-check.json et evidence/attacks/results.json, ainsi que les captures et textes associés.',
'Dépôt d’origine':'Nous attribuons DVWA à ses auteurs : https://github.com/digininja/DVWA. Notre rapport s’appuie sur les sources locales du projet et les preuves enregistrées.'
}
for p in d.paragraphs:
 text=p.text
 for prefix,new in replacements.items():
  if text.startswith(prefix): text=new; break
 text=text.replace('/home/modou/dvwa-study','/home/modou/groupe12').replace('dvwa-study/','groupe12/')
 text=text.replace('Résultat du contrôle de démarrage du 14 septembre','Résultat de notre contrôle de démarrage du 15 septembre').replace('Contrôle HTTP rejoué le 14 septembre 2026','Notre contrôle HTTP du 15 septembre 2026')
 if text!=p.text:
  p.text=text
# Insérer le chapitre avant les annexes, sans modifier la V1.
anchor=next(p._p for p in d.paragraphs if p.text.startswith('Annexe A.'))
added=[]
def para(text,style=None):
 p=d.add_paragraph(text,style); added.append(p._p); return p

def heading(text,level=1): return para(text,'Heading '+str(level))
def page():
 p=d.add_page_break(); added.append(p._p)
entries=json.loads((R/'evidence/attacks/results.json').read_text())
titles=['Injection SQL','Injection SQL aveugle','Injection de commandes','XSS réfléchie','XSS stockée','XSS DOM','Requête intersite CSRF','Inclusion de fichiers','Téléversement de fichier exécutable','Contournement d’autorisation','Contrôle d’accès et cookie falsifié','Essais de mots de passe','Identifiants de session faibles','Contournement du parcours CAPTCHA','Réutilisation d’un nonce CSP','Jeton calculé côté client','Redirection ouverte','Cryptographie XOR','Recomposition de blocs ECB','API : comparaison v1 et v2','API : modification et relecture','API : injection de commandes']
limits={
'T06':'Nous conservons les deux captures, mais le contrôle automatique du marqueur a échoué. Nous ne confirmons pas l’exécution JavaScript pour cet essai et devons revoir le contexte DOM et la charge.',
'T13':'Nous avons ouvert le module. Le test a été interrompu par une navigation qui a détruit le contexte JavaScript ; nous ne disposons pas d’une série de jetons permettant de conclure.',
'T18a':'Nous voyons un message décodé sur la capture. Toutefois, la vérification attendue du script a échoué : nous conservons ce sous-test à reprendre sans annoncer sa validation complète.',
'T19b':'Nous recevons level=0 dans la réponse de modification, mais la relecture renvoie level=1, comme avant le test. Nous ne confirmons donc pas une élévation de privilèges persistante, malgré le champ success du journal.',
'T19c':'Nous ne disposons d’aucune capture pour cet essai et le journal le classe à reprendre. Nous ne confirmons pas d’exécution de commande via l’API.',
'T12':'Nous limitons notre conclusion à trois essais, dont le dernier utilise le mot de passe connu. Nous n’en déduisons pas une absence générale de limitation à grande échelle.',
'T07':'Nous avons utilisé deux origines locales distinctes par leur port, mais appartenant au même site. Nous ne généralisons pas ce résultat à tous les contextes intersites.',
'T14':'Nous confirmons le changement par la reconnexion enregistrée avec le nouveau mot de passe. Nous n’avons résolu aucun CAPTCHA externe ; le parcours normal avec reCAPTCHA reste à valider.',
'T18b':'Nous confirmons la recomposition acceptée dans ce sous-exercice. Nous ne présentons pas ce résultat comme une rupture de l’algorithme AES ni comme une attaque par oracle complète.'}
def verdict(e):
 if e['id']=='T19b': return 'Modification persistante non confirmée'
 return 'Effet confirmé' if e.get('success') else 'À reprendre'
heading('9. Nos expérimentations et leurs preuves')
para('Nous présentons les essais enregistrés le 14 septembre 2026 sur notre instance de référence, port 4280. Les pages témoins CSRF et de redirection utilisent le port 4290. Nous avons regroupé les 39 captures originales selon les identifiants des essais et leur ordre dans le journal.')
para('Nous utilisons le niveau low par défaut. T15 (CSP) et T18b (ECB) utilisent medium. T10 et T11 utilisent le compte ordinaire gordonb ; les autres connexions DVWA utilisent admin. Pour les routes API, nous distinguons la version appelée du niveau sélectionné dans l’interface.')
para('Nous recensons 22 sous-tests sur 19 modules : 17 présentent un effet confirmé par les vérifications enregistrées, quatre restent à reprendre et un présente une modification non confirmée à la relecture. Nous n’assimilons pas le seul statut « exécuté » à une exploitation réussie.')
t=d.add_table(rows=1,cols=3); added.append(t._tbl); t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Essai','Objet','Notre bilan']): c.text=v
for e,title in zip(entries,titles):
 for c,v in zip(t.add_row().cells,[e['id'],title,verdict(e)]): c.text=v
fig=max([int(m.group(1)) for p in d.paragraphs if (m:=re.match(r'Figure (\d+) —',p.text))] or [0])
for index,(e,title) in enumerate(zip(entries,titles),1):
 page(); heading(f'9.{index}. {e["id"]} — {title}',2)
 para(f'Notre essai du {e["date"]} • Module : {e["module"]} • Bilan : {verdict(e)}.')
 if e.get('payload'):
  para('Nous avons soumis l’entrée suivante ou réalisé cette séquence :')
  para(e['payload'],'Code')
 if e.get('observed'):
  obs=e['observed']
  if e['id']=='T02': obs='La condition vraie renvoie HTTP 200 et « User ID exists ». La condition fausse renvoie HTTP 404 et « MISSING ». Nous obtenons donc deux réponses distinctes selon la condition SQL.'
  para('Nous relevons le résultat suivant dans nos preuves : '+obs)
 if e['id'] in limits: para(limits[e['id']])
 elif e.get('success'): para('Nous retenons ce résultat pour le scénario enregistré. Nous n’avons pas de preuve de rejeu après correctif personnel dans cette série de captures.')
 for j,cap in enumerate(e['captures']):
  if j: page()
  fig+=1
  para(f'Nous présentons ci-dessous la figure {fig}, correspondant à l’étape « {cap["label"].replace("-"," ")} » de notre essai {e["id"]}.')
  f=R/'docs/captures/attacks'/cap['file']; w,h=struct.unpack('>II',f.read_bytes()[16:24])
  p=para(''); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
  p.add_run().add_picture(str(f),width=Inches(min(6.65,7.0*w/h)))
  para(f'Figure {fig} — {e["id"]} : {title} — {cap["label"].replace("-"," ")}', 'FigureCaption')
  para('Source : docs/captures/attacks/'+cap['file'],'Caption')
page(); heading('10. Notre bilan et la suite du travail')
para('Nous avons mis en service deux instances séparées, conservé leurs données et vérifié leur disponibilité. Nous avons intégré les 39 captures d’essais disponibles pour relier nos observations à des preuves identifiables.')
para('Nos essais confirment plusieurs effets sur les injections, les traitements de fichiers, les autorisations et certaines fonctions côté navigateur. Nous distinguons ces résultats des quatre sous-tests à reprendre et de la modification API dont la persistance n’est pas confirmée.')
para('Nous devons compléter T06, T13, T18a et T19c, puis approfondir T19b en vérifiant la logique de mise à jour et l’état relu. Nous compléterons ensuite nos comparaisons par des correctifs identifiés dans Git, un rejeu des attaques et des contrôles des usages légitimes. Les captures présentes ne constituent pas une preuve de correction personnelle.')
for el in added: anchor.addprevious(el)
# Éviter les anciens résultats de champs : LibreOffice actualisera les index.
for s in d.sections:
 for p in s.header.paragraphs:
  if 'GROUPE' in p.text: p.text='GROUPE 12 • Notre étude DVWA • Rapport V2'
out=R/'GROUPE-12/Rapport-DVWA-V2.docx'; d.save(out)
print(out)
print('Captures attaques intégrées :',sum(len(e['captures']) for e in entries))
